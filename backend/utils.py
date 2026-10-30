import aiofiles
import urllib
import mistune
import os
import re
import shutil
import subprocess
import sys
import tempfile


async def write_to_file(filename: str, text: str) -> None:
    """Asynchronously write text to a file in UTF-8 encoding.

    Args:
        filename: The filename to write to.
        text: The text to write.
    """
    # Ensure text is a string
    if not isinstance(text, str):
        text = str(text)

    # Convert to UTF-8, replacing any problematic characters
    text_utf8 = text.encode('utf-8', errors='replace').decode('utf-8')

    async with aiofiles.open(filename, "w", encoding='utf-8') as file:
        await file.write(text_utf8)


async def write_text_to_md(text: str, filename: str = "") -> str:
    """Writes text to a Markdown file and returns the file path.

    Args:
        text: Text to write to the Markdown file.
        filename: Base name for the file.

    Returns:
        The file path of the generated Markdown file.
    """
    import uuid

    safe_name = (filename or "").strip()[:60] or f"report-{uuid.uuid4().hex[:12]}"
    safe_name = safe_name.replace("/", "-").replace("\\", "-")
    os.makedirs("outputs", exist_ok=True)
    file_path = f"outputs/{safe_name}.md"
    await write_to_file(file_path, text)
    return urllib.parse.quote(file_path)


def _strip_front_matter(text: str) -> str:
    """Drop the YAML metadata block the agent puts at the top of a report."""
    if text.lstrip().startswith("---"):
        match = re.match(r"^\s*---\s*\n.*?\n---\s*\n", text, flags=re.S)
        if match:
            return text[match.end():]
    return text


def _mark_blockquotes(html: str) -> str:
    """Tag paragraphs inside <blockquote> so they can skip first-line indent."""
    def repl(match):
        inner = match.group(1).replace('<p>', '<p style="margin-left: 20px;">')
        return f"<blockquote>{inner}</blockquote>"

    return re.sub(r"<blockquote>(.*?)</blockquote>", repl, html, flags=re.S)


def _preprocess_images_for_pdf(text: str) -> str:
    """Convert web image URLs to absolute file paths for PDF generation."""
    base_path = os.path.abspath(".")

    def replace_image_url(match):
        alt_text = match.group(1)
        url = match.group(2)

        if url.startswith("/outputs/"):
            abs_path = os.path.join(base_path, url.lstrip("/"))
            return f"![{alt_text}]({abs_path})"
        return match.group(0)

    pattern = r'!\[([^\]]*)\]\((/outputs/[^)]+)\)'
    return re.sub(pattern, replace_image_url, text)


def _find_headless_browser():
    """Return a Chromium-based browser executable usable for PDF printing."""
    if sys.platform != "win32":
        return None
    candidates = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    ]
    for path in candidates:
        if os.path.exists(path):
            return path
    return None


def _render_pdf_with_browser(html_body: str, pdf_path: str, css_path: str) -> bool:
    """Render HTML to PDF with headless Edge/Chrome (no GTK/WeasyPrint needed)."""
    browser = _find_headless_browser()
    if not browser:
        return False

    try:
        css = ""
        if os.path.exists(css_path):
            with open(css_path, "r", encoding="utf-8") as handle:
                css = handle.read()

        # Chromium wants file:// URLs, not raw Windows paths.
        html_body = re.sub(
            r'src="([A-Za-z]:[\\/][^"]+)"',
            lambda m: 'src="file:///' + m.group(1).replace("\\", "/") + '"',
            html_body,
        )

        document = (
            '<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">'
            f"<style>{css}</style></head><body><div class=\"markdown-body\">"
            f"{html_body}</div></body></html>"
        )

        tmp_dir = tempfile.mkdtemp(prefix="gptr-pdf-")
        html_path = os.path.join(tmp_dir, "report.html")
        profile_dir = os.path.join(tmp_dir, "profile")
        with open(html_path, "w", encoding="utf-8") as handle:
            handle.write(document)

        if os.path.exists(pdf_path):
            os.remove(pdf_path)

        creationflags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        subprocess.run(
            [
                browser,
                "--headless=new",
                "--disable-gpu",
                "--no-sandbox",
                f"--user-data-dir={profile_dir}",
                "--no-pdf-header-footer",
                "--allow-file-access-from-files",
                f"--print-to-pdf={os.path.abspath(pdf_path)}",
                html_path,
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            timeout=180,
            creationflags=creationflags,
        )

        ok = os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 0
        shutil.rmtree(tmp_dir, ignore_errors=True)
        return ok
    except Exception as e:
        print(f"Error in browser PDF rendering: {e}")
        return False


async def write_md_to_pdf(text: str, filename: str = "") -> str:
    """Converts Markdown text to a PDF file and returns the file path.

    Tries headless Edge/Chrome first (Windows has neither GTK nor WeasyPrint
    dependencies), then falls back to WeasyPrint via md2pdf elsewhere.
    """
    import uuid

    safe_name = (filename or "").strip()[:60] or f"report-{uuid.uuid4().hex[:12]}"
    safe_name = safe_name.replace("/", "-").replace("\\", "-")
    os.makedirs("outputs", exist_ok=True)
    file_path = f"outputs/{safe_name}.pdf"

    current_dir = os.path.dirname(os.path.abspath(__file__))
    css_path = os.path.join(current_dir, "styles", "pdf_styles.css")

    body_text = _strip_front_matter(text)
    processed_text = _preprocess_images_for_pdf(body_text)

    try:
        html_body = mistune.html(processed_text)
        if _render_pdf_with_browser(html_body, file_path, css_path):
            print(f"Report written to {file_path}")
            return urllib.parse.quote(file_path)
    except Exception as e:
        print(f"Error in converting Markdown to PDF (browser): {e}")

    try:
        from md2pdf.core import md2pdf
        md2pdf(
            file_path,
            raw=processed_text,
            css=css_path,
            base_url=os.path.abspath("."),
        )
        print(f"Report written to {file_path}")
    except Exception as e:
        print(f"Error in converting Markdown to PDF: {e}")
        return ""

    return urllib.parse.quote(file_path)


def _apply_docx_style(doc) -> None:
    """Apply the requested Chinese document formatting.

    - body: SimSun (宋体), 小四 = 12pt
    - latin/digits: Times New Roman
    - line spacing: exactly 20pt
    - margins: left 3cm, right 2.5cm
    """
    from docx.shared import Pt, Cm, RGBColor
    from docx.oxml.ns import qn

    def set_style_fonts(style, size_pt, bold=None):
        style.font.name = "Times New Roman"
        style.font.size = Pt(size_pt)
        style.font.color.rgb = RGBColor(0, 0, 0)
        if bold is not None:
            style.font.bold = bold
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), "宋体")
        rfonts.set(qn("w:ascii"), "Times New Roman")
        rfonts.set(qn("w:hAnsi"), "Times New Roman")
        style.paragraph_format.line_spacing = Pt(20)

    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)

    set_style_fonts(doc.styles["Normal"], 12)
    heading_sizes = {1: 16, 2: 14, 3: 12, 4: 12, 5: 12, 6: 12}
    for level, size in heading_sizes.items():
        try:
            set_style_fonts(doc.styles[f"Heading {level}"], size, bold=True)
        except KeyError:
            continue

    def is_heading(paragraph):
        try:
            return (paragraph.style.name or "").startswith(("Heading", "Title"))
        except Exception:
            return False

    def style_paragraph(paragraph, in_table=False):
        paragraph.paragraph_format.line_spacing = Pt(20)
        heading = is_heading(paragraph)
        style_name = paragraph.style.name or ""
        is_list = style_name.startswith("List")
        left_indent = paragraph.paragraph_format.left_indent
        is_indented_block = left_indent is not None and left_indent > 0
        has_text = bool(paragraph.text.strip())

        # 整段正文首行缩进 2 字符（12pt × 2）；标题、列表、引用、表格不缩进
        if in_table or heading or is_list or is_indented_block or not has_text:
            paragraph.paragraph_format.first_line_indent = Pt(0)
        else:
            paragraph.paragraph_format.first_line_indent = Pt(24)

        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            if not heading:
                run.font.size = Pt(12)
            rpr = run._element.get_or_add_rPr()
            rfonts = rpr.get_or_add_rFonts()
            rfonts.set(qn("w:eastAsia"), "宋体")
            rfonts.set(qn("w:ascii"), "Times New Roman")
            rfonts.set(qn("w:hAnsi"), "Times New Roman")

    for paragraph in doc.paragraphs:
        style_paragraph(paragraph)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    style_paragraph(paragraph, in_table=True)


async def write_md_to_word(text: str, filename: str = "") -> str:
    """Converts Markdown text to a DOCX file and returns the file path."""
    import uuid

    safe_name = (filename or "").strip()[:60] or f"report-{uuid.uuid4().hex[:12]}"
    safe_name = safe_name.replace("/", "-").replace("\\", "-")
    os.makedirs("outputs", exist_ok=True)
    file_path = f"outputs/{safe_name}.docx"

    try:
        from docx import Document
        from htmldocx import HtmlToDocx

        body_text = _strip_front_matter(text)
        html = _mark_blockquotes(mistune.html(body_text))
        doc = Document()
        HtmlToDocx().add_html_to_document(html, doc)
        _apply_docx_style(doc)
        doc.save(file_path)

        print(f"Report written to {file_path}")

        encoded_file_path = urllib.parse.quote(file_path)
        return encoded_file_path

    except Exception as e:
        print(f"Error in converting Markdown to DOCX: {e}")
        return ""